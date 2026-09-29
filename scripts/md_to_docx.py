"""Render a Markdown document as a styled Word file.

Written for `docs/solution-overview.md`, which is the document a reviewer is asked to read
before the walkthrough — and which is easier to circulate as `.docx` than as Markdown. The
conversion is deliberately narrow: it handles exactly the Markdown that document uses
(ATX headings, GFM pipe tables, fenced code, block quotes, bullet and numbered lists, and
inline bold/italic/code/links) and nothing else. A general-purpose converter would be
pandoc, which is not installed here.

Two choices worth knowing:

* every `#` heading starts a new page, so the Parts of the document are also its pages —
  the assumptions and the not-achieved sections are each asked for as "a page";
* only `http(s)` links become real hyperlinks. A relative link to a sibling Markdown file
  would be dead in a Word document circulated on its own, so it renders as plain text.

`python-docx` is not a dependency of this project — it is not needed to build, test or run
anything. Run this with an interpreter that has it:

    python3 scripts/md_to_docx.py docs/solution-overview.md docs/Solution-Overview.docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

MONO = "Menlo"
CODE_GREY = "F2F4F7"
HEADER_BLUE = "1461C4"
LINK_BLUE = RGBColor(0x14, 0x61, 0xC4)
MUTED = RGBColor(0x5D, 0x6B, 0x7C)

#: `**bold**`, `` `code` ``, `[text](target)`, `*italic*`. Matching is leftmost-first, and
#: bold and italic recurse into their own content, because the document nests them: a bold
#: span holding inline code (`**no `ALARM_API_*` field**`) is common in it, and a
#: non-recursive pass renders the backticks as literal text. The italic arm requires
#: non-space on both sides of the content so that a lone `*` inside an identifier such as
#: `ALARM_API_*` cannot open an emphasis run that never closes.
INLINE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|`(?P<code>[^`]+)`"
    r"|\[(?P<text>[^\]]+)\]\((?P<href>[^)]+)\)"
    r"|\*(?P<italic>\S(?:[^*]*\S)?)\*"
)


def _shade(element: Any, fill: str) -> None:
    """Paint a cell or paragraph background. python-docx has no API for this."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _hyperlink(paragraph: Any, text: str, href: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        href,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), HEADER_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(colour)
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def write_inline(
    paragraph: Any,
    text: str,
    *,
    size: float | None = None,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """Add `text` to `paragraph`, honouring inline markup.

    `bold` and `italic` are inherited by every run produced, which is how nesting works:
    the bold arm calls back into this function with `bold=True`.
    """

    def emphasise(run: Any) -> None:
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)

    def plain(chunk: str) -> None:
        if chunk:
            emphasise(paragraph.add_run(chunk))

    cursor = 0
    for match in INLINE.finditer(text):
        plain(text[cursor : match.start()])
        cursor = match.end()
        if match.group("bold"):
            write_inline(paragraph, match.group("bold"), size=size, bold=True, italic=italic)
        elif match.group("code"):
            run = paragraph.add_run(match.group("code"))
            emphasise(run)
            run.font.name = MONO
            run.font.size = Pt((size or 10.5) - 1.5)
        elif match.group("text"):
            href = match.group("href")
            label = match.group("text")
            if href.startswith("http"):
                # The hyperlink run is built in raw XML, so its label cannot carry nested
                # markup — the backticks around a filename are dropped rather than printed.
                _hyperlink(paragraph, label.replace("`", ""), href)
            else:
                write_inline(paragraph, label, size=size, bold=bold, italic=italic)
        elif match.group("italic"):
            write_inline(paragraph, match.group("italic"), size=size, bold=bold, italic=True)
    plain(text[cursor:])


def _row_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_divider(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


class Renderer:
    """Walks the Markdown line by line and appends to a `Document`."""

    def __init__(self, document: Any) -> None:
        self.doc = document
        self.pages = 0

    def heading(self, level: int, text: str) -> None:
        if level == 1:
            self.pages += 1
            if self.pages > 1:
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        paragraph = self.doc.add_heading(level=min(level, 4))
        paragraph.paragraph_format.space_before = Pt(16 if level > 1 else 4)
        paragraph.paragraph_format.space_after = Pt(6)
        write_inline(paragraph, text)
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)

    def code(self, lines: list[str]) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.15)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(10)
        _shade(paragraph._p.get_or_add_pPr(), CODE_GREY)
        for index, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.name = MONO
            run.font.size = Pt(8.5)
            if index < len(lines) - 1:
                run.add_break()

    def quote(self, lines: list[str]) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(10)
        write_inline(paragraph, " ".join(lines))
        for run in paragraph.runs:
            run.italic = True

    def table(self, rows: list[list[str]]) -> None:
        header, *body = rows
        table = self.doc.add_table(rows=len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for column, label in enumerate(header):
            cell = table.cell(0, column)
            cell.text = ""
            _shade(cell._tc.get_or_add_tcPr(), CODE_GREY)
            write_inline(cell.paragraphs[0], label, size=9)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for index, row in enumerate(body, start=1):
            for column in range(len(header)):
                cell = table.cell(index, column)
                cell.text = ""
                value = row[column] if column < len(row) else ""
                write_inline(cell.paragraphs[0], value, size=9)
        for row_obj in table.rows:
            for cell in row_obj.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def bullet(self, text: str, *, ordered: bool, indent: int = 0) -> None:
        style = "List Number" if ordered else "List Bullet"
        paragraph = self.doc.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Inches(0.3 + 0.25 * indent)
        paragraph.paragraph_format.space_after = Pt(3)
        write_inline(paragraph, text)

    def body(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(9)
        write_inline(paragraph, text)


def render(markdown: str, document: Any) -> None:
    renderer = Renderer(document)
    lines = markdown.splitlines()
    index = 0
    paragraph: list[str] = []

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            renderer.body(" ".join(paragraph))
            paragraph = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped or stripped == "---":
            flush()
            index += 1
            continue

        if stripped.startswith("```"):
            flush()
            index += 1
            block: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            renderer.code(block)
            index += 1
            continue

        heading = re.match(r"(#{1,6})\s+(.*)", stripped)
        if heading:
            flush()
            renderer.heading(len(heading.group(1)), heading.group(2))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and _is_divider(lines[index + 1]):
            flush()
            rows = [_row_cells(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_row_cells(lines[index]))
                index += 1
            renderer.table(rows)
            continue

        if stripped.startswith(">"):
            flush()
            quoted: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quoted.append(lines[index].strip().lstrip(">").strip())
                index += 1
            renderer.quote([q for q in quoted if q])
            continue

        item = re.match(r"(\s*)(?:([-*])|(\d+)\.)\s+(.*)", line)
        if item:
            flush()
            text = item.group(4)
            index += 1
            # A list item's continuation lines are indented; fold them into the item.
            while index < len(lines) and re.match(r"\s{3,}\S", lines[index]):
                if re.match(r"\s*(?:[-*]|\d+\.)\s+", lines[index]):
                    break
                text += " " + lines[index].strip()
                index += 1
            renderer.bullet(
                text,
                ordered=item.group(3) is not None,
                indent=len(item.group(1)) // 2,
            )
            continue

        paragraph.append(stripped)
        index += 1

    flush()


def build(source: Path, target: Path) -> None:
    document = Document()

    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    lines = source.read_text().splitlines()
    title = lines[0].lstrip("# ").strip()
    subtitle = next((line.strip() for line in lines[1:6] if line.strip()), "")

    heading = document.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(HEADER_BLUE)
    if subtitle:
        lede = document.add_paragraph()
        write_inline(lede, subtitle)
        for run in lede.runs:
            run.font.color.rgb = MUTED
            run.bold = False

    document.core_properties.title = title
    document.core_properties.comments = f"Generated from {source.name} by scripts/md_to_docx.py"

    render("\n".join(lines[1:]), document)
    document.save(target)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(f"usage: {Path(argv[0]).name} <source.md> <target.docx>", file=sys.stderr)
        return 2
    source, target = Path(argv[1]), Path(argv[2])
    build(source, target)
    print(f"Wrote {target} from {source}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
